import os
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

import requests
import urllib3
from docx import Document, document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from dotenv import load_dotenv, dotenv_values
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ==========================================================
# CONFIGURATION
# ==========================================================

# Determine the application folder first.
if getattr(sys, "frozen", False):
    # Running from a PyInstaller executable.
    SCRIPT_FOLDER = Path(sys.executable).resolve().parent
else:
    # Running from a Python script.
    SCRIPT_FOLDER = Path(__file__).resolve().parent

# Always look for .env beside the script or executable.
ENV_FILE = SCRIPT_FOLDER / ".env"

# Capture environment variables BEFORE loading .env.
# This allows us to distinguish real environment variables
# from values that python-dotenv adds to os.environ.
existing_environment = {
    "SMARTSHEET_TOKEN": os.environ.get("SMARTSHEET_TOKEN"),
    "SMARTSHEET_SHEET_ID": os.environ.get("SMARTSHEET_SHEET_ID"),
}

# Read the .env file separately for source detection.
if ENV_FILE.exists():
    env_file_values = dotenv_values(ENV_FILE)
else:
    env_file_values = {}

# Load .env values, but do not overwrite environment variables
# that were already present when the application started.
load_dotenv(
    dotenv_path=ENV_FILE,
    override=False,
)

# Retrieve the effective configuration values.
TOKEN = os.getenv("SMARTSHEET_TOKEN")
SHEET_ID = os.getenv("SMARTSHEET_SHEET_ID")


def get_configuration_source(variable_name: str) -> str:
    """
    Identify the source of the effective configuration value.

    Existing operating-system environment variables take priority
    because load_dotenv() is called with override=False.
    """
    existing_value = existing_environment.get(variable_name)
    dotenv_value = env_file_values.get(variable_name)

    if existing_value:
        if dotenv_value:
            return "Environment variable (takes priority over .env)"
        return "Environment variable"

    if dotenv_value:
        return f".env file: {ENV_FILE}"

    return "Not found"


token_source = get_configuration_source("SMARTSHEET_TOKEN")
sheet_source = get_configuration_source("SMARTSHEET_SHEET_ID")


def configuration_error(
    variable_name: str,
    description: str,
    example_value: str,
    source: str,
) -> None:
    """
    Display a clean configuration error and exit without a traceback.
    """
    print()
    print("=" * 70)
    print(f"ERROR: {variable_name} is not configured")
    print("=" * 70)
    print()
    print(f"This application requires {description}.")
    print()
    print("Configure it using one of these methods:")
    print()
    print("Option 1: Add it to the .env file")
    print(f"  File: {ENV_FILE}")
    print(f"  Entry: {variable_name}={example_value}")
    print()
    print("Option 2: Set a Windows environment variable")
    print(f"  set {variable_name}={example_value}")
    print()
    print("Current configuration:")
    print(f"  Application folder: {SCRIPT_FOLDER}")
    print(f"  .env file found   : {ENV_FILE.exists()}")
    print(f"  Detected source   : {source}")
    print()
    print("=" * 70)

    sys.exit(1)


# Validate configuration.
if not TOKEN:
    configuration_error(
        variable_name="SMARTSHEET_TOKEN",
        description="a Smartsheet API token",
        example_value="your_token_here",
        source=token_source,
    )

if not SHEET_ID:
    configuration_error(
        variable_name="SMARTSHEET_SHEET_ID",
        description="a Smartsheet Sheet ID",
        example_value="your_sheet_id_here",
        source=sheet_source,
    )

# Report the detected configuration sources without displaying
# the token or Sheet ID values.
print(f"SMARTSHEET_TOKEN loaded from: {token_source}")
print(f"SMARTSHEET_SHEET_ID loaded from: {sheet_source}")

# Create the output folder.
OUTPUT_FOLDER = SCRIPT_FOLDER / "output"
OUTPUT_FOLDER.mkdir(
    parents=True,
    exist_ok=True,
)

# These filenames remain unchanged each day.
# Your Copilot agent should reference these files.
DOCX_FILE = OUTPUT_FOLDER / "project_summary.docx"
XLSX_FILE = OUTPUT_FOLDER / "project_data.xlsx"

# Set this to None to include every row in Word summary sections.
# The complete inventories always include every record.
SUMMARY_SECTION_LIMIT: Optional[int] = 50  # noqa: UP045

# Workbook table names cannot contain spaces.
EXCEL_TABLE_STYLE = "TableStyleMedium2"


# ==========================================================
# DOWNLOAD SMARTSHEET
# ==========================================================

url = f"https://api.smartsheet.com/2.0/sheets/{SHEET_ID}?include=rowPermalink"

headers = {
    "Authorization": f"Bearer {TOKEN}",
    "Accept": "application/json",
}

print("Downloading Smartsheet data...")

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

response = requests.get(url, headers=headers, timeout=120, verify=False)
response.raise_for_status()
data = response.json()

sheet_name = data.get("name", "Unnamed Smartsheet")

print(f"Sheet Name : {sheet_name}")
print(f"Rows       : {len(data.get('rows', []))}")


# ==========================================================
# COLUMN MAP
# ==========================================================

column_map = {
    column["id"]: column.get("title", str(column["id"]))
    for column in data.get("columns", [])
}


# ==========================================================
# GENERAL HELPERS
# ==========================================================


def clean_value(value: Any) -> str:
    """
    Convert None and other values to safe display strings.
    """
    if value is None:
        return ""

    if isinstance(value, bool):
        return "Yes" if value else "No"

    return str(value).strip()


def get_cell_value(cell: dict) -> Any:
    """
    Prefer displayValue when it exists. Otherwise use value.
    """
    if cell.get("displayValue") is not None:
        return cell.get("displayValue")

    return cell.get("value")


def build_row_dict(row: dict) -> dict:
    """
    Convert a Smartsheet row's cells into:
        {"Column Title": value}
    """
    row_data = {}

    for cell in row.get("cells", []):
        column_name = column_map.get(cell.get("columnId"))

        if not column_name:
            continue

        row_data[column_name] = get_cell_value(cell)

    return row_data


def parse_date(value: Any) -> Optional[date]:  # noqa: UP045
    """
    Parse common Smartsheet date formats.
    """
    if value in (None, ""):
        return None

    value_text = str(value).strip()

    formats = [
        "%Y-%m-%d",
        "%m/%d/%Y",
        "%m/%d/%y",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%dT%H:%M:%SZ",
        "%Y-%m-%dT%H:%M:%S.%fZ",
    ]

    for date_format in formats:
        try:
            return datetime.strptime(value_text, date_format).date()  # noqa: DTZ007
        except ValueError:
            continue

    # Handle ISO strings containing timezone information.
    try:
        normalized = value_text.replace("Z", "+00:00")
        return datetime.fromisoformat(normalized).date()
    except ValueError:
        return None


def safe_float(value: Any) -> float:
    """
    Convert a numeric or percentage value to a float.
    """
    if value in (None, ""):
        return 0.0

    try:
        cleaned = str(value).replace("%", "").replace(",", "").strip()
        return float(cleaned)
    except (TypeError, ValueError):
        return 0.0


def bool_value(value: Any) -> bool:
    """
    Convert common truthy values to a Boolean.
    """
    if value is None:
        return False

    normalized = str(value).strip().lower()

    return normalized in {
        "yes",
        "y",
        "true",
        "1",
        "red",
        "high",
        "at risk",
    }


def normalize_percentage(value: float) -> float:
    """
    Smartsheet may return percentages as either:
        0.50
    or:
        50

    Convert values between 0 and 1 to 0 through 100.
    """
    if 0 < value <= 1:
        return value * 100

    return value


def display_number(value: Any) -> str:
    """
    Display whole-number floats without a trailing .0.
    """
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value)


def calculate_risk_reason(
    overdue: bool,
    at_risk: bool,
    schedule_variance: float,
    assigned_to: Any,
) -> str:
    """
    Explain why the calculated risk score was assigned.
    """
    reasons = []

    if overdue:
        reasons.append("Overdue")

    if at_risk:
        reasons.append("Marked at risk")

    if schedule_variance < -20:
        reasons.append(
            f"Behind expected completion by "
            f"{abs(schedule_variance):g} percentage points"
        )

    if not assigned_to:
        reasons.append("No owner assigned")

    if not reasons:
        return "No calculated risk indicators"

    return "; ".join(reasons)


def limit_items(items: list[dict]) -> list[dict]:
    """
    Limit only the highlight sections.

    Complete inventories are never limited.
    """
    if SUMMARY_SECTION_LIMIT is None:
        return items

    return items[:SUMMARY_SECTION_LIMIT]


# ==========================================================
# PROCESS SMARTSHEET ROWS
# ==========================================================

today = date.today()  # noqa: DTZ011

tasks: list[dict] = []
parent_rows: list[dict] = []
phase_rows: list[dict] = []

for row in data.get("rows", []):
    row_data = build_row_dict(row)
    row_number = row.get("rowNumber")

    # task_name = row_data.get("Task Name")
    task_name = clean_value(row_data.get("Task Name"))

    if not task_name:
        continue

    start_date = parse_date(row_data.get("Start Date"))
    end_date = parse_date(row_data.get("End Date"))

    percent_complete = normalize_percentage(safe_float(row_data.get("% Complete")))

    expected_percent = normalize_percentage(safe_float(row_data.get("% Expected")))

    assigned_to = row_data.get("Assigned To")
    status = row_data.get("Status")
    at_risk = bool_value(row_data.get("At Risk"))
    predecessors = row_data.get("Predecessors")
    notes = row_data.get("Notes")
    parent = row_data.get("Parent")
    children = row_data.get("Children")
    ancestors = row_data.get("Ancestors")
    duration = row_data.get("Duration")
    level = row_data.get("Level")
    task_id = row_data.get("Task ID")

    is_parent = bool_value(row_data.get("Is Parent Row"))
    is_phase = bool_value(row_data.get("Is Phase Row"))

    row_permalink = row.get("rowPermalink") or row.get("permalink") or ""

    # ------------------------------------------------------
    # SCHEDULE CALCULATIONS
    # ------------------------------------------------------

    schedule_variance = percent_complete - expected_percent

    overdue = False
    days_overdue = 0

    if end_date and percent_complete < 100:
        overdue = end_date < today

        if overdue:
            days_overdue = (today - end_date).days

    days_until_due: Optional[int] = None  # noqa: UP045
    due_next_7_days = False
    due_next_30_days = False

    if end_date:
        days_until_due = (end_date - today).days
        due_next_7_days = 0 <= days_until_due <= 7
        due_next_30_days = 0 <= days_until_due <= 30

    # ------------------------------------------------------
    # RISK SCORE
    # ------------------------------------------------------

    risk_score = 0

    if overdue:
        risk_score += 50

    if at_risk:
        risk_score += 30

    if schedule_variance < -20:
        risk_score += 20

    if not assigned_to:
        risk_score += 10

    if risk_score >= 70:
        risk_level = "High"
    elif risk_score >= 40:
        risk_level = "Medium"
    else:
        risk_level = "Low"

    risk_reason = calculate_risk_reason(
        overdue=overdue,
        at_risk=at_risk,
        schedule_variance=schedule_variance,
        assigned_to=assigned_to,
    )

    # ------------------------------------------------------
    # RECORD TYPE
    # ------------------------------------------------------

    if is_phase:
        record_type = "Phase"
    elif is_parent:
        record_type = "Parent"
    else:
        record_type = "Task"

    # ------------------------------------------------------
    # COMPLETE RECORD
    # ------------------------------------------------------

    task_record = {
        "record_type": record_type,
        "row_number": row_number,
        "task_id": task_id,
        "task_name": task_name,
        "assigned_to": assigned_to,
        "status": status,
        "at_risk": at_risk,
        "percent_complete": percent_complete,
        "expected_percent": expected_percent,
        "schedule_variance": schedule_variance,
        "start_date": start_date.isoformat() if start_date else "",
        "end_date": end_date.isoformat() if end_date else "",
        "duration": duration,
        "predecessors": predecessors,
        "notes": notes,
        "parent": parent,
        "children": children,
        "ancestors": ancestors,
        "level": level,
        "is_parent_row": is_parent,
        "is_phase_row": is_phase,
        "days_until_due": days_until_due,
        "overdue": overdue,
        "days_overdue": days_overdue,
        "due_next_7_days": due_next_7_days,
        "due_next_30_days": due_next_30_days,
        "risk_score": risk_score,
        "risk_level": risk_level,
        "risk_reason": risk_reason,
        "row_id": row.get("id"),
        "row_permalink": row_permalink,
    }

    if is_phase:
        phase_rows.append(task_record)
    elif is_parent:
        parent_rows.append(task_record)
    else:
        tasks.append(task_record)


# ==========================================================
# BUILD LEAF-TASK VIEWS
# ==========================================================

overdue_tasks = [task for task in tasks if task["overdue"]]

at_risk_tasks = [task for task in tasks if task["at_risk"]]

coming_due_7_days = [task for task in tasks if task["due_next_7_days"]]

coming_due_30_days = [task for task in tasks if task["due_next_30_days"]]

behind_schedule_tasks = [task for task in tasks if task["schedule_variance"] < -20]

unassigned_tasks = [task for task in tasks if not task["assigned_to"]]

high_risk_tasks = [task for task in tasks if task["risk_level"] == "High"]

overdue_parent_rows = [task for task in parent_rows if task["overdue"]]

at_risk_parent_rows = [task for task in parent_rows if task["at_risk"]]

overdue_phases = [task for task in phase_rows if task["overdue"]]

at_risk_phases = [task for task in phase_rows if task["at_risk"]]


# ==========================================================
# SORT VIEWS
# ==========================================================

overdue_tasks.sort(
    key=lambda item: item["days_overdue"],
    reverse=True,
)

at_risk_tasks.sort(
    key=lambda item: item["risk_score"],
    reverse=True,
)

high_risk_tasks.sort(
    key=lambda item: item["risk_score"],
    reverse=True,
)

coming_due_7_days.sort(
    key=lambda item: (
        item["days_until_due"] if item["days_until_due"] is not None else 999999
    )
)

coming_due_30_days.sort(
    key=lambda item: (
        item["days_until_due"] if item["days_until_due"] is not None else 999999
    )
)

behind_schedule_tasks.sort(key=lambda item: item["schedule_variance"])

tasks.sort(
    key=lambda item: (
        clean_value(item.get("wbs")),
        clean_value(item.get("task_name")),
    )
)

parent_rows.sort(
    key=lambda item: (
        clean_value(item.get("wbs")),
        clean_value(item.get("task_name")),
    )
)

phase_rows.sort(
    key=lambda item: (
        clean_value(item.get("wbs")),
        clean_value(item.get("task_name")),
    )
)


# ==========================================================
# SUMMARY METRICS
# ==========================================================

summary_metrics = [
    ("Leaf Tasks", len(tasks)),
    ("Parent Rows", len(parent_rows)),
    ("Phase Rows", len(phase_rows)),
    ("High-Risk Tasks", len(high_risk_tasks)),
    ("Overdue Tasks", len(overdue_tasks)),
    ("Tasks Marked At Risk", len(at_risk_tasks)),
    ("Tasks Due Within 7 Days", len(coming_due_7_days)),
    ("Tasks Due Within 30 Days", len(coming_due_30_days)),
    ("Tasks Behind Schedule", len(behind_schedule_tasks)),
    ("Unassigned Tasks", len(unassigned_tasks)),
    ("At-Risk Parent Rows", len(at_risk_parent_rows)),
    ("Overdue Parent Rows", len(overdue_parent_rows)),
    ("At-Risk Phases", len(at_risk_phases)),
    ("Overdue Phases", len(overdue_phases)),
]


# ==========================================================
# WORD HELPERS
# ==========================================================


def set_cell_shading(cell, fill: str) -> None:
    """
    Apply a background color to a Word table cell.
    """
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell_properties.append(shading)


def set_cell_text(
    cell,
    value: Any,
    bold: bool = False,
    color: Optional[str] = None,  # noqa: UP045
) -> None:
    """
    Replace cell content with formatted text.
    """
    cell.text = ""

    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(clean_value(value))
    run.bold = bold

    if color:
        run.font.color.rgb = RGBColor.from_string(color)

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_word_hyperlink(
    paragraph,
    text: str,
    url_value: str,
) -> None:
    """
    Add a true clickable hyperlink to a Word paragraph.
    """
    if not url_value:
        return

    relationship_id = paragraph.part.relate_to(
        url_value,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    new_run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    run_properties.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)

    new_run.append(run_properties)

    text_element = OxmlElement("w:t")
    text_element.text = text

    new_run.append(text_element)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_word_field(
    document: Document,
    label: str,
    value: Any,
) -> None:
    """
    Add a labeled field to a Word task record.
    """
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)

    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True

    paragraph.add_run(clean_value(value))


def add_word_link_field(
    document: Document,
    label: str,
    url_value: str,
) -> None:
    """
    Add a labeled clickable link to Word.
    """
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)

    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True

    if url_value:
        add_word_hyperlink(
            paragraph,
            "Open Smartsheet row",
            url_value,
        )
    else:
        paragraph.add_run("")


def add_word_task_table(
    document: Document,
    title: str,
    records: list[dict],
    limit_records: bool = True,
) -> None:
    """
    Add a compact highlight table to Word.
    """
    document.add_heading(title, level=2)

    display_records = limit_items(records) if limit_records else records

    if not display_records:
        document.add_paragraph("None found.")
        return

    table = document.add_table(rows=1, cols=9)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headings = [
        "Row #",
        "Task",
        "Owner",
        "End Date",
        "Risk",
        "Risk Reason",
        "Status",
        "% Complete",
        "Link",
    ]

    for index, heading in enumerate(headings):
        set_cell_text(
            table.rows[0].cells[index],
            heading,
            bold=True,
            color="FFFFFF",
        )
        set_cell_shading(
            table.rows[0].cells[index],
            "1F4E78",
        )

    for record in display_records:
        cells = table.add_row().cells

        values = [
            record.get("row_number"),
            record.get("task_name"),
            record.get("assigned_to"),
            record.get("end_date"),
            record.get("risk_level"),
            record.get("risk_reason"),
            record.get("status"),
            display_number(record.get("percent_complete")),
        ]

        for index, value in enumerate(values):
            set_cell_text(cells[index], value)

        risk_level = record.get("risk_level")

        if risk_level == "High":
            set_cell_shading(cells[4], "F4CCCC")
        elif risk_level == "Medium":
            set_cell_shading(cells[4], "FCE5CD")
        else:
            set_cell_shading(cells[4], "D9EAD3")

        link_paragraph = cells[8].paragraphs[0]

        if record.get("row_permalink"):
            add_word_hyperlink(
                link_paragraph,
                "Link",
                record["row_permalink"],
            )

    if SUMMARY_SECTION_LIMIT is not None and len(records) > len(display_records):
        document.add_paragraph(
            f"Showing {len(display_records)} of "
            f"{len(records)} records. "
            f"See project_data.xlsx for the complete data set."
        )

    document.add_paragraph()


def add_complete_word_inventory(
    document: Document,
    title: str,
    records: list[dict],
    heading_level: int = 2,
) -> None:
    """
    Add every record as searchable heading and field-value text.

    This format is intentionally used instead of one extremely
    large Word table so task names, owners, notes, hierarchy,
    dates, and dependencies remain explicit searchable text.
    """
    document.add_heading(title, level=1)

    if not records:
        document.add_paragraph("None found.")
        return

    for index, record in enumerate(records, start=1):
        task_name = clean_value(record.get("task_name"))
        task_id = clean_value(record.get("task_id"))

        if task_name and task_id:
            heading_text = f"{task_name} [{task_id}]"
        elif task_name:
            heading_text = task_name
        elif task_id:
            heading_text = f"Task {task_id}"
        else:
            heading_text = f"Record {index}"

        document.add_heading(
            heading_text,
            level=heading_level,
        )

        add_word_field(
            document,
            "Record Type",
            record.get("record_type"),
        )
        add_word_field(
            document,
            "Row #",
            record.get("row_number"),
        )
        add_word_field(
            document,
            "WBS",
            record.get("wbs"),
        )
        add_word_field(
            document,
            "Task Name",
            record.get("task_name"),
        )
        add_word_field(
            document,
            "Owner",
            record.get("assigned_to"),
        )
        add_word_field(
            document,
            "Status",
            record.get("status"),
        )
        add_word_field(
            document,
            "At Risk",
            record.get("at_risk"),
        )
        add_word_field(
            document,
            "Risk Level",
            record.get("risk_level"),
        )
        add_word_field(
            document,
            "Risk Score",
            record.get("risk_score"),
        )
        add_word_field(
            document,
            "Risk Reason",
            record.get("risk_reason"),
        )
        add_word_field(
            document,
            "Percent Complete",
            display_number(record.get("percent_complete")),
        )
        add_word_field(
            document,
            "Expected Percent Complete",
            display_number(record.get("expected_percent")),
        )
        add_word_field(
            document,
            "Schedule Variance",
            display_number(record.get("schedule_variance")),
        )
        add_word_field(
            document,
            "Start Date",
            record.get("start_date"),
        )
        add_word_field(
            document,
            "End Date",
            record.get("end_date"),
        )
        add_word_field(
            document,
            "Duration",
            record.get("duration"),
        )
        add_word_field(
            document,
            "Days Until Due",
            record.get("days_until_due"),
        )
        add_word_field(
            document,
            "Overdue",
            record.get("overdue"),
        )
        add_word_field(
            document,
            "Days Overdue",
            record.get("days_overdue"),
        )
        add_word_field(
            document,
            "Due Within 7 Days",
            record.get("due_next_7_days"),
        )
        add_word_field(
            document,
            "Due Within 30 Days",
            record.get("due_next_30_days"),
        )
        add_word_field(
            document,
            "Parent",
            record.get("parent"),
        )
        add_word_field(
            document,
            "Ancestors",
            record.get("ancestors"),
        )
        add_word_field(
            document,
            "Children",
            record.get("children"),
        )
        add_word_field(
            document,
            "Predecessors",
            record.get("predecessors"),
        )
        add_word_field(
            document,
            "Level",
            record.get("level"),
        )
        add_word_field(
            document,
            "Notes",
            record.get("notes"),
        )
        add_word_field(
            document,
            "Smartsheet Row ID",
            record.get("row_id"),
        )
        add_word_link_field(
            document,
            "Smartsheet Link",
            clean_value(record.get("row_permalink")),
        )

        separator = document.add_paragraph()
        separator.paragraph_format.space_after = Pt(4)
        separator.add_run("─" * 80)


# ==========================================================
# CREATE WORD DOCUMENT
# ==========================================================


def create_word_document() -> None:
    document = Document()

    # Standard margins for narrative sections.
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = document.styles

    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)

    styles["Title"].font.name = "Aptos Display"
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 2"].font.name = "Aptos Display"

    title = document.add_heading(
        "Project Health Summary",
        level=0,
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(sheet_name).bold = True

    generated_paragraph = document.add_paragraph()
    generated_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    generated_paragraph.add_run(
        f"Generated: {datetime.now().astimezone().strftime('%Y-%m-%d %H:%M:%S %Z (%z)')}"
    )

    document.add_paragraph(
        "This document contains project-health summaries and "
        "highlighted records requiring attention."
    )

    # ------------------------------------------------------
    # EXECUTIVE SUMMARY
    # ------------------------------------------------------

    document.add_heading("Executive Summary", level=1)

    executive_text = (
        f"The export contains {len(tasks)} leaf tasks, "
        f"{len(parent_rows)} parent rows, and "
        f"{len(phase_rows)} phase rows. "
        f"There are {len(overdue_tasks)} overdue leaf tasks, "
        f"{len(high_risk_tasks)} high-risk leaf tasks, "
        f"{len(at_risk_tasks)} leaf tasks explicitly marked at risk, "
        f"{len(behind_schedule_tasks)} leaf tasks more than "
        f"20 percentage points behind expected completion, and "
        f"{len(unassigned_tasks)} leaf tasks without an owner. "
        f"{len(coming_due_7_days)} leaf tasks are due within "
        f"the next 7 days, and {len(coming_due_30_days)} are due "
        f"within the next 30 days."
    )

    document.add_paragraph(executive_text)

    document.add_paragraph(
        "Calculated risk levels use the following rules: "
        "50 points for overdue work, 30 points when the Smartsheet "
        "At Risk field is true, 20 points when completion is more "
        "than 20 percentage points behind expectation, and "
        "10 points when no owner is assigned. "
        "High risk is 70 or more points, medium risk is 40 through "
        "69 points, and low risk is below 40 points."
    )

    # ------------------------------------------------------
    # PORTFOLIO METRICS
    # ------------------------------------------------------

    document.add_heading("Portfolio Metrics", level=1)

    metrics_table = document.add_table(rows=1, cols=2)
    metrics_table.style = "Table Grid"
    metrics_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    set_cell_text(
        metrics_table.rows[0].cells[0],
        "Metric",
        bold=True,
        color="FFFFFF",
    )
    set_cell_text(
        metrics_table.rows[0].cells[1],
        "Count",
        bold=True,
        color="FFFFFF",
    )
    set_cell_shading(
        metrics_table.rows[0].cells[0],
        "1F4E78",
    )
    set_cell_shading(
        metrics_table.rows[0].cells[1],
        "1F4E78",
    )

    for metric_name, metric_value in summary_metrics:
        cells = metrics_table.add_row().cells
        set_cell_text(cells[0], metric_name)
        set_cell_text(cells[1], metric_value)

    document.add_page_break()

    # ------------------------------------------------------
    # HIGHLIGHT SECTIONS
    # ------------------------------------------------------

    add_word_task_table(
        document,
        "High-Risk Leaf Tasks",
        high_risk_tasks,
    )

    add_word_task_table(
        document,
        "Overdue Leaf Tasks",
        overdue_tasks,
    )

    add_word_task_table(
        document,
        "Leaf Tasks Marked At Risk",
        at_risk_tasks,
    )

    add_word_task_table(
        document,
        "Leaf Tasks Due Within 7 Days",
        coming_due_7_days,
    )

    add_word_task_table(
        document,
        "Leaf Tasks Due Within 30 Days",
        coming_due_30_days,
    )

    add_word_task_table(
        document,
        "Leaf Tasks Behind Schedule",
        behind_schedule_tasks,
    )

    add_word_task_table(
        document,
        "Unassigned Leaf Tasks",
        unassigned_tasks,
    )

    add_word_task_table(
        document,
        "At-Risk Parent Rows",
        at_risk_parent_rows,
    )

    add_word_task_table(
        document,
        "Overdue Parent Rows",
        overdue_parent_rows,
    )

    add_word_task_table(
        document,
        "At-Risk Phases",
        at_risk_phases,
    )

    add_word_task_table(
        document,
        "Overdue Phases",
        overdue_phases,
    )

    document.save(DOCX_FILE)

    # ------------------------------------------------------
    # FULL SEARCHABLE INVENTORIES
    # ------------------------------------------------------

    """document.add_page_break()

    add_complete_word_inventory(
        document,
        "Complete Leaf Task Inventory",
        tasks,
    )

    document.add_page_break()

    add_complete_word_inventory(
        document,
        "Complete Parent Row Inventory",
        parent_rows,
    )

    document.add_page_break()

    add_complete_word_inventory(
        document,
        "Complete Phase Inventory",
        phase_rows,
    ) """


# ==========================================================
# EXCEL HELPERS
# ==========================================================

EXCEL_FIELDS = [
    ("Record Type", "record_type"),
    ("WBS", "wbs"),
    ("Row #", "row_number"),
    ("Task Name", "task_name"),
    ("Owner", "assigned_to"),
    ("Status", "status"),
    ("At Risk", "at_risk"),
    ("Risk Level", "risk_level"),
    ("Risk Score", "risk_score"),
    ("Risk Reason", "risk_reason"),
    ("% Complete", "percent_complete"),
    ("% Expected", "expected_percent"),
    ("Schedule Variance", "schedule_variance"),
    ("Start Date", "start_date"),
    ("End Date", "end_date"),
    ("Duration", "duration"),
    ("Days Until Due", "days_until_due"),
    ("Overdue", "overdue"),
    ("Days Overdue", "days_overdue"),
    ("Due Next 7 Days", "due_next_7_days"),
    ("Due Next 30 Days", "due_next_30_days"),
    ("Parent", "parent"),
    ("Ancestors", "ancestors"),
    ("Children", "children"),
    ("Predecessors", "predecessors"),
    ("Level", "level"),
    ("Notes", "notes"),
    ("Is Parent Row", "is_parent_row"),
    ("Is Phase Row", "is_phase_row"),
    ("Smartsheet Row ID", "row_id"),
    ("Smartsheet Link", "row_permalink"),
]


def safe_excel_value(value: Any) -> Any:
    """
    Preserve numbers and Booleans while converting None to blank.
    """
    if value is None:
        return ""

    if isinstance(value, (int, float, bool)):
        return value

    return str(value)


def create_excel_record_sheet(
    workbook: Workbook,
    sheet_title: str,
    records: list[dict],
    table_name: str,
) -> None:
    """
    Create a structured Excel worksheet from record dictionaries.
    """
    worksheet = workbook.create_sheet(sheet_title)

    headers = [display_name for display_name, _ in EXCEL_FIELDS]

    worksheet.append(headers)

    for record in records:
        worksheet.append(
            [safe_excel_value(record.get(field_name)) for _, field_name in EXCEL_FIELDS]
        )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    worksheet.freeze_panes = "A2"
    # worksheet.auto_filter.ref = worksheet.dimensions

    if records:
        table_reference = f"A1:{get_column_letter(len(headers))}{len(records) + 1}"

        excel_table = Table(
            displayName=table_name,
            ref=table_reference,
        )

        table_style = TableStyleInfo(
            name=EXCEL_TABLE_STYLE,
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        excel_table.tableStyleInfo = table_style
        worksheet.add_table(excel_table)

    header_index = {header: index + 1 for index, header in enumerate(headers)}

    risk_column = header_index["Risk Level"]
    link_column = header_index["Smartsheet Link"]
    notes_column = header_index["Notes"]
    risk_reason_column = header_index["Risk Reason"]

    for row_number in range(2, worksheet.max_row + 1):
        risk_cell = worksheet.cell(
            row=row_number,
            column=risk_column,
        )

        if risk_cell.value == "High":
            risk_cell.fill = PatternFill(
                fill_type="solid",
                fgColor="F4CCCC",
            )
        elif risk_cell.value == "Medium":
            risk_cell.fill = PatternFill(
                fill_type="solid",
                fgColor="FCE5CD",
            )
        elif risk_cell.value == "Low":
            risk_cell.fill = PatternFill(
                fill_type="solid",
                fgColor="D9EAD3",
            )

        link_cell = worksheet.cell(
            row=row_number,
            column=link_column,
        )

        if link_cell.value:
            link_cell.hyperlink = str(link_cell.value)
            link_cell.value = "Open row"
            link_cell.style = "Hyperlink"

        worksheet.cell(
            row=row_number,
            column=notes_column,
        ).alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

        worksheet.cell(
            row=row_number,
            column=risk_reason_column,
        ).alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

    widths = {
        "A": 14,
        "B": 13,
        "C": 14,
        "D": 42,
        "E": 25,
        "F": 20,
        "G": 12,
        "H": 13,
        "I": 12,
        "J": 42,
        "K": 13,
        "L": 13,
        "M": 18,
        "N": 13,
        "O": 13,
        "P": 13,
        "Q": 15,
        "R": 11,
        "S": 13,
        "T": 16,
        "U": 17,
        "V": 35,
        "W": 40,
        "X": 35,
        "Y": 30,
        "Z": 10,
        "AA": 55,
        "AB": 14,
        "AC": 13,
        "AD": 20,
        "AE": 14,
    }

    for column_letter, width in widths.items():
        worksheet.column_dimensions[column_letter].width = width

    for row in worksheet.iter_rows():
        for cell in row:
            if cell.row > 1:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=cell.column
                    in {
                        4,
                        10,
                        22,
                        23,
                        24,
                        25,
                        27,
                    },
                )


def create_excel_summary_sheet(
    workbook: Workbook,
) -> None:
    """
    Create the workbook summary worksheet.
    """
    worksheet = workbook.active
    worksheet.title = "Summary"

    worksheet["A1"] = "Project Health Summary"
    worksheet["A1"].font = Font(
        size=18,
        bold=True,
        color="FFFFFF",
    )
    worksheet["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    worksheet.merge_cells("A1:D1")

    worksheet["A2"] = "Sheet"
    worksheet["B2"] = sheet_name

    worksheet["A3"] = "Generated"
    worksheet["B3"] = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %Z (%z)")

    worksheet["A5"] = "Metric"
    worksheet["B5"] = "Count"

    for row in worksheet["A5:B5"]:
        for cell in row:
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )

            cell.font = Font(
                color="FFFFFF",
                bold=True,
            )

    for row_number, (metric_name, metric_value) in enumerate(
        summary_metrics,
        start=6,
    ):
        worksheet.cell(
            row=row_number,
            column=1,
            value=metric_name,
        )
        worksheet.cell(
            row=row_number,
            column=2,
            value=metric_value,
        )

    summary_end_row = 5 + len(summary_metrics)

    summary_table = Table(
        displayName="SummaryMetrics",
        ref=f"A5:B{summary_end_row}",
    )

    summary_table.tableStyleInfo = TableStyleInfo(
        name=EXCEL_TABLE_STYLE,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    worksheet.add_table(summary_table)

    executive_row = summary_end_row + 3

    worksheet.cell(
        row=executive_row,
        column=1,
        value="Executive Summary",
    )
    worksheet.cell(
        row=executive_row,
        column=1,
    ).font = Font(
        bold=True,
        size=14,
        color="FFFFFF",
    )
    worksheet.cell(
        row=executive_row,
        column=1,
    ).fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    worksheet.merge_cells(
        start_row=executive_row,
        start_column=1,
        end_row=executive_row,
        end_column=4,
    )

    executive_text = (
        f"The export contains {len(tasks)} leaf tasks, "
        f"{len(parent_rows)} parent rows, and "
        f"{len(phase_rows)} phase rows. "
        f"There are {len(overdue_tasks)} overdue leaf tasks, "
        f"{len(high_risk_tasks)} high-risk leaf tasks, "
        f"{len(at_risk_tasks)} leaf tasks marked at risk, "
        f"{len(behind_schedule_tasks)} leaf tasks behind schedule, "
        f"and {len(unassigned_tasks)} unassigned leaf tasks."
    )

    worksheet.cell(
        row=executive_row + 1,
        column=1,
        value=executive_text,
    )

    worksheet.merge_cells(
        start_row=executive_row + 1,
        start_column=1,
        end_row=executive_row + 3,
        end_column=4,
    )

    worksheet.cell(
        row=executive_row + 1,
        column=1,
    ).alignment = Alignment(
        vertical="top",
        wrap_text=True,
    )

    worksheet.column_dimensions["A"].width = 36
    worksheet.column_dimensions["B"].width = 18
    worksheet.column_dimensions["C"].width = 18
    worksheet.column_dimensions["D"].width = 18

    worksheet.freeze_panes = "A5"


def create_excel_workbook() -> None:
    workbook = Workbook()

    create_excel_summary_sheet(workbook)

    create_excel_record_sheet(
        workbook,
        sheet_title="Tasks",
        records=tasks,
        table_name="LeafTasksTable",
    )

    create_excel_record_sheet(
        workbook,
        sheet_title="Parent Rows",
        records=parent_rows,
        table_name="ParentRowsTable",
    )

    create_excel_record_sheet(
        workbook,
        sheet_title="Phase Rows",
        records=phase_rows,
        table_name="PhaseRowsTable",
    )

    create_excel_record_sheet(
        workbook,
        sheet_title="High Risk",
        records=high_risk_tasks,
        table_name="HighRiskTable",
    )

    create_excel_record_sheet(
        workbook,
        sheet_title="Overdue",
        records=overdue_tasks,
        table_name="OverdueTasksTable",
    )

    create_excel_record_sheet(
        workbook,
        sheet_title="At Risk",
        records=at_risk_tasks,
        table_name="AtRiskTasksTable",
    )

    create_excel_record_sheet(
        workbook,
        sheet_title="Due 7 Days",
        records=coming_due_7_days,
        table_name="DueSevenDaysTable",
    )

    create_excel_record_sheet(
        workbook,
        sheet_title="Due 30 Days",
        records=coming_due_30_days,
        table_name="DueThirtyDaysTable",
    )

    create_excel_record_sheet(
        workbook,
        sheet_title="Behind Schedule",
        records=behind_schedule_tasks,
        table_name="BehindScheduleTable",
    )

    create_excel_record_sheet(
        workbook,
        sheet_title="Unassigned",
        records=unassigned_tasks,
        table_name="UnassignedTasksTable",
    )

    workbook.save(XLSX_FILE)


# ==========================================================
# GENERATE COPILOT FILES
# ==========================================================

print()
print("Creating Copilot-readable Word document...")
create_word_document()

print("Creating Copilot-readable Excel workbook...")
create_excel_workbook()


# ==========================================================
# RESULTS
# ==========================================================

print()
print("=" * 70)
print("EXPORT COMPLETE")
print("=" * 70)

print(f"Sheet             : {sheet_name}")
print(f"Leaf Tasks        : {len(tasks)}")
print(f"Parent Rows       : {len(parent_rows)}")
print(f"Phase Rows        : {len(phase_rows)}")
print(f"High Risk         : {len(high_risk_tasks)}")
print(f"Overdue           : {len(overdue_tasks)}")
print(f"Marked At Risk    : {len(at_risk_tasks)}")
print(f"Due Within 7 Days : {len(coming_due_7_days)}")
print(f"Due Within 30 Days: {len(coming_due_30_days)}")
print(f"Behind Schedule   : {len(behind_schedule_tasks)}")
print(f"Unassigned        : {len(unassigned_tasks)}")

print()
print(f"Word Knowledge File : {DOCX_FILE}")
print(f"Excel Data File      : {XLSX_FILE}")

print()
print("Add both fixed-name files individually to the Copilot agent:")
print(f"  1. {DOCX_FILE.name}")
print(f"  2. {XLSX_FILE.name}")
